import io
import os
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("PB_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("PB_DATABASE_URL", f"sqlite:///{tmp_path / 'test.db'}")
    monkeypatch.setenv("PB_AI_PROVIDER", "mock")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    from app import config, db, security
    from app.drafting import ai
    config.settings = config.Settings()
    security.reset_crypto()
    ai.set_provider(None)
    db.init_engine(config.settings.database_url)
    db.create_all()
    from fastapi.testclient import TestClient
    from app.main import app
    with TestClient(app) as c:
        c.headers.update({"x-requested-with": "policybuilder"})
        yield c
    ai.set_provider(None)


def register(c, email="owner@northstar.example", ws="Northstar workspace"):
    r = c.post("/api/auth/register", json={"email": email, "password": "correct-horse-battery", "display_name": "Test Owner",
                                           "workspace_name": ws})
    assert r.status_code == 200, r.text
    return r


def new_project(c, name="Northstar policies", lang="en", mode="standard"):
    r = c.post("/api/projects", json={"name": name, "mode": mode, "doc_language": lang})
    assert r.status_code == 200, r.text
    return r.json()["id"]


def answer(c, pid, code, value=None, status="answered"):
    r = c.put(f"/api/projects/{pid}/answers/{code}", json={"value": value, "status": status})
    assert r.status_code == 200, (code, r.text)
    return r.json()


# ---------------------------------------------------------------- synthetic organizations (fictional)
SOFTWARE_QC = {  # (A) 35-person Montreal software firm using AWS / Microsoft 365
    "ORG-01": "Northstar Digital Services Inc.", "ORG-01B": "Northstar", "ORG-02": "software_saas",
    "ORG-03": "Northstar builds and operates a B2B scheduling SaaS platform for clinics and retailers.",
    "ORG-04": "11-50", "ORG-05": ["QC", "ON"], "ORG-06": "private", "ORG-08": "scratch",
    "ORG-09": ["governance", "client_assurance"], "ORG-12": "developing", "ORG-13": ["it_manager"],
    "ORG-14": "The scheduling platform and its customer database", "DATA-01": ["employee", "customer_pi", "client_confidential"],
    "TECH-01": ["m365", "aws"], "TECH-02": "hybrid", "TECH-02B": "no", "ORG-15": "recommended",
}
FOOD_OT = {  # (B) 600-person Quebec food manufacturer with plants and outsourced IT
    "ORG-01": "Laiterie Boréale ltée", "ORG-02": "manufacturing_ot", "ORG-03": "Dairy processing plants supplying grocery chains.",
    "ORG-04": "251-1000", "ORG-05": ["QC"], "ORG-06": "private", "ORG-08": "some", "ORG-09": ["governance", "incidents"],
    "ORG-12": "informal", "ORG-13": ["msp", "hr"], "ORG-14": "Production lines, cold-chain monitoring and ERP",
    "DATA-01": ["employee", "ot_data", "financial"], "TECH-01": ["m365", "onprem"], "TECH-02": "onsite", "TECH-02B": "limited",
    "ORG-15": "recommended", "TECH-03": ["ot", "iot"],
}
HEALTH = {  # (C) 2,000-person multi-jurisdiction healthcare services network
    "ORG-01": "Meridian Care Network", "ORG-02": "healthcare", "ORG-03": "Outpatient clinics and home-care services.",
    "ORG-04": "1001-5000", "ORG-05": ["QC", "ON", "US"], "ORG-06": "healthcare", "ORG-08": "update",
    "ORG-09": ["governance", "privacy", "audit"], "ORG-12": "documented",
    "ORG-13": ["ciso", "security_team", "privacy_officer", "risk_committee", "internal_audit", "legal", "hr"],
    "ORG-13B": "Chief Information Security Officer", "ORG-14": "Electronic health records and scheduling",
    "DATA-01": ["employee", "customer_pi", "health", "financial"], "TECH-01": ["azure", "m365", "onprem"],
    "TECH-02": "hybrid", "TECH-02B": "limited", "ORG-15": "recommended",
}
UNIVERSITY = {  # (D) university with research and student records
    "ORG-01": "Université du Littoral (fictive)", "ORG-02": "education_research", "ORG-03": "Teaching and research university.",
    "ORG-04": "1001-5000", "ORG-05": ["QC"], "ORG-06": "education", "ORG-08": "some", "ORG-09": ["governance", "privacy"],
    "ORG-12": "developing", "ORG-13": ["ciso", "privacy_officer", "legal"], "ORG-14": "Student information system and research data",
    "DATA-01": ["employee", "students", "research", "ip"], "TECH-01": ["m365", "onprem", "azure"], "TECH-02": "hybrid",
    "TECH-02B": "yes", "ORG-15": "recommended",
}
ORGS = {"software_qc": SOFTWARE_QC, "food_ot": FOOD_OT, "health": HEALTH, "university": UNIVERSITY}


def apply_answers(c, pid, answers):
    for k, v in answers.items():
        answer(c, pid, k, v)


def make_docx(paragraphs, title=None) -> bytes:
    import docx
    d = docx.Document()
    if title:
        d.core_properties.title = title
        d.add_heading(title, level=1)
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
