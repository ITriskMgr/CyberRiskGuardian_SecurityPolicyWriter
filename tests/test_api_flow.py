"""End-to-end vertical slice through the API (AT-01, AT-04, AT-06, AT-07, AT-08, AT-11, AT-12, AT-16)."""
import io
import zipfile

import docx

from conftest import SOFTWARE_QC, answer, apply_answers, make_docx, new_project, register

CORE6 = ["EISP", "AUP", "IAM", "CLASS", "IR", "BACKUP"]


def _read_docx(data: bytes) -> str:
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    for s in d.sections:
        parts.extend(p.text for p in s.header.paragraphs)
    return "\n".join(parts)


def test_full_vertical_slice(client):
    c = client
    register(c)
    pid = new_project(c)
    apply_answers(c, pid, SOFTWARE_QC)
    prof = c.get(f"/api/projects/{pid}/profile").json()
    assert prof["progress"]["minimum_complete"]
    assert prof["org_name"] == "Northstar"
    assert any(o["code"] == "GP-NO-PRIVACY-OFFICER" for o in prof["observations"])

    # evidence: an existing access policy (AT-04)
    data = make_docx(["All access to company systems must use multi-factor authentication for all users.",
                      "Critical vulnerabilities must be remediated within 14 days.",
                      "Information is classified as Public, Internal, Confidential or Restricted."], title="Access Control Policy v2")
    r = c.post(f"/api/projects/{pid}/sources", files={"file": ("access-policy.docx", data)}, data={"category": "existing_policy"})
    assert r.status_code == 200, r.text
    assert r.json()["facts"] >= 3
    facts = c.get(f"/api/projects/{pid}/facts").json()
    mfa = [f for f in facts if f["fact_type"] == "DEC-IAM-MFA-SCOPE"][0]
    assert mfa["value"] == "all" and mfa["state"] == "proposed" and mfa["locator"].startswith("§")
    assert "multi-factor" in mfa["quote"]
    assert c.post(f"/api/facts/{mfa['id']}/decision", json={"action": "confirm"}).status_code == 200
    cls = [f for f in facts if f["fact_type"] == "DATA-02B"][0]
    c.post(f"/api/facts/{cls['id']}/decision", json={"action": "confirm"})

    # catalogue and selection (EISP + 5)
    cat = c.get(f"/api/projects/{pid}/catalogue").json()
    rec = {i["code"] for i in cat["items"] if i["recommended"]}
    assert {"EISP", "IAM", "IR", "SDLC", "PRIV"} <= rec
    assert "OT" not in rec
    assert c.put(f"/api/projects/{pid}/selection", json={"codes": CORE6}).status_code == 200
    dec = c.get(f"/api/projects/{pid}/decisions").json()
    assert {g["policy"] for g in dec} >= {"EISP", "IAM", "IR"}
    answer(c, pid, "DEC-IAM-LEAVER", "same_day")

    # generation
    r = c.post(f"/api/projects/{pid}/generate", json={})
    assert r.status_code == 200, r.text
    pols = c.get(f"/api/projects/{pid}/policies").json()
    assert len(pols) == 6 and all(p["version"] == 1 for p in pols)
    iam = [p for p in pols if p["code"] == "IAM"][0]
    body = c.get(f"/api/policies/{iam['id']}").json()["body"]
    texts = " ".join(cl["text"] for s in body["sections"] for cl in s["clauses"])
    assert "Multi-factor authentication must be used by all users" in texts          # confirmed document fact
    assert "no later than the end of their last working day" in texts                # decision
    assert "[DECISION: role(s) that approve access requests]" in texts               # open decision, not invented
    eisp = [p for p in pols if p["code"] == "EISP"][0]
    eb = c.get(f"/api/policies/{eisp['id']}").json()["body"]
    assert "[APPROVER:" in " ".join(cl["text"] for s in eb["sections"] for cl in s["clauses"])   # AT-08
    assert any(d["type"] == "APPROVER" for d in eb["open_decisions"])
    issues = c.get(f"/api/projects/{pid}/issues").json()
    assert any(i["kind"] == "unresolved_placeholders" for i in issues)

    # single docx export (AT-08: exportable with placeholders)
    r = c.post(f"/api/projects/{pid}/exports", json={"kind": "docx", "document_id": eisp["id"], "lang": "en"})
    assert r.status_code == 200, r.text
    dl = c.get(r.json()["download"])
    assert dl.status_code == 200
    txt = _read_docx(dl.content)
    assert "DRAFT — NOT APPROVED" in txt and "NOR-EISP-001" in txt and "Enterprise Information Security Policy" in txt
    d = docx.Document(io.BytesIO(dl.content))
    assert d.core_properties.identifier.startswith("NOR-EISP-001|" + eisp["id"])
    assert any(p.style.name == "Heading 1" for p in d.paragraphs)

    # ZIP package (AT-07)
    r = c.post(f"/api/projects/{pid}/exports", json={"kind": "zip"})
    z = zipfile.ZipFile(io.BytesIO(c.get(r.json()["download"]).content))
    names = z.namelist()
    assert sum(n.endswith("_DRAFT.docx") for n in names) == 6
    assert any("/01_Governance/NOR-EISP-001" in n for n in names)
    assert any("/02_Policies/NOR-POL-IAM-001" in n for n in names)
    assert any(n.endswith("traceability.csv") for n in names)
    for n in names:
        if n.endswith(".docx"):
            docx.Document(io.BytesIO(z.read(n)))  # every file opens

    # change an answer -> stale -> regenerate creates v2, v1 intact (AT-11, AT-16)
    answer(c, pid, "TECH-02B", "yes")
    pols = c.get(f"/api/projects/{pid}/policies").json()
    assert [p for p in pols if p["code"] == "AUP"][0]["stale"]
    c.post(f"/api/projects/{pid}/generate", json={"codes": ["AUP"]})
    aup = [p for p in c.get(f"/api/projects/{pid}/policies").json() if p["code"] == "AUP"][0]
    assert aup["version"] == 2 and not aup["stale"]
    v1 = c.get(f"/api/policies/{aup['id']}?version=1").json()["body"]
    v2 = c.get(f"/api/policies/{aup['id']}?version=2").json()["body"]
    t1 = " ".join(cl["text"] for s in v1["sections"] for cl in s["clauses"])
    t2 = " ".join(cl["text"] for s in v2["sections"] for cl in s["clauses"])
    assert "Personal devices must not be used" in t1 and "Personal devices must not be used" not in t2
    assert "remote work and mobile device policy" in t2
    assert any(i["kind"] == "broken_cross_reference" for i in c.get(f"/api/projects/{pid}/issues").json())

    # manual clause edit -> new version, earlier kept
    r = c.patch(f"/api/policies/{aup['id']}/clauses/AUP-G01", json={"text": "Users must use Northstar technology for work.", "lang": "en"})
    assert r.status_code == 200 and r.json()["version_no"] == 3
    assert c.get(f"/api/policies/{aup['id']}?version=2").json()["body"]["sections"]

    # project reload
    proj = c.get(f"/api/projects/{pid}").json()
    assert proj["workflow_state"] == "DRAFT_GENERATED" and proj["generated"] == 6


def test_single_policy_only(client):  # AT-06
    c = client
    register(c)
    pid = new_project(c)
    apply_answers(c, pid, SOFTWARE_QC)
    c.put(f"/api/projects/{pid}/selection", json={"codes": ["IAM"]})
    c.post(f"/api/projects/{pid}/generate", json={})
    pols = c.get(f"/api/projects/{pid}/policies").json()
    assert [p["code"] for p in pols] == ["IAM"]
    body = c.get(f"/api/policies/{pols[0]['id']}").json()["body"]
    assert body["parent_policy_id"] == "NOR-EISP-001"
    assert any(i["kind"] == "missing_parent" for i in c.get(f"/api/projects/{pid}/issues").json())
    r = c.post(f"/api/projects/{pid}/exports", json={"kind": "zip"})
    z = zipfile.ZipFile(io.BytesIO(c.get(r.json()["download"]).content))
    assert sum(n.endswith("_DRAFT.docx") for n in z.namelist()) == 1


def test_bilingual_output(client):  # AT-12
    c = client
    register(c)
    pid = new_project(c, lang="both")
    apply_answers(c, pid, {k: v for k, v in SOFTWARE_QC.items()})
    answer(c, pid, "ORG-10", "both")
    c.put(f"/api/projects/{pid}/selection", json={"codes": ["EISP", "IAM"]})
    c.post(f"/api/projects/{pid}/generate", json={})
    iam = [p for p in c.get(f"/api/projects/{pid}/policies").json() if p["code"] == "IAM"][0]
    assert iam["languages"] == ["en", "fr"]
    en = c.get(f"/api/policies/{iam['id']}?lang=en").json()["body"]
    fr = c.get(f"/api/policies/{iam['id']}?lang=fr").json()["body"]
    ids = lambda b: [cl["clause_id"] for s in b["sections"] for cl in s["clauses"]]  # noqa: E731
    assert ids(en) == ids(fr)
    assert fr["title"] == "Politique de gestion des identités et des accès"
    assert fr["sections"][0]["heading"] == "Objet"
    r = c.post(f"/api/projects/{pid}/exports", json={"kind": "docx", "document_id": iam["id"], "lang": "fr"})
    txt = _read_docx(c.get(r.json()["download"]).content)
    assert "ÉBAUCHE — NON APPROUVÉE" in txt and "Contrôle du document" in txt
    r = c.post(f"/api/projects/{pid}/exports", json={"kind": "zip"})
    names = zipfile.ZipFile(io.BytesIO(c.get(r.json()["download"]).content)).namelist()
    assert any("/EN/" in n for n in names) and any("/FR/" in n for n in names)


def test_project_delete_removes_files(client, tmp_path):  # AT-13
    c = client
    register(c)
    pid = new_project(c)
    apply_answers(c, pid, SOFTWARE_QC)
    c.post(f"/api/projects/{pid}/sources", files={"file": ("notes.txt", b"Backups are taken daily.")})
    from app import config
    folder = config.settings.blob_dir
    assert any(folder.rglob("*.bin"))
    assert c.delete(f"/api/projects/{pid}").status_code == 200
    assert not any(folder.rglob("*.bin"))
    assert c.get(f"/api/projects/{pid}").status_code == 404
