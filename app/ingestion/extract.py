"""Safe file identification and text extraction with locators (spec §6.2, §14.1).

Uploaded files are untrusted: the type is sniffed from content, sizes and archive expansion are bounded,
no macros or embedded objects are executed, and text that looks like instructions to an AI is flagged.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field

MAX_ZIP_UNCOMPRESSED = 200 * 1024 * 1024
MAX_ZIP_RATIO = 150
MAX_ZIP_ENTRIES = 5000
SEG_CHARS = 1500

ALLOWED = {"docx", "pdf", "txt", "md", "csv", "xlsx", "image"}

INJECTION_PATTERNS = [
    r"ignore\s+(all\s+|any\s+)?(the\s+)?(previous|prior|above|earlier)\s+(instructions|prompts?|rules)",
    r"disregard\s+(all\s+|the\s+)?(previous|prior|above|system)\s+(instructions|prompt)",
    r"you\s+are\s+now\s+(a|an|the)\b", r"\bsystem\s+prompt\b", r"reveal\s+(your|the)\s+(instructions|prompt|secrets?|api\s*key)",
    r"\bact\s+as\s+(a|an)\s+(different|new)\b", r"do\s+not\s+follow\s+the\s+user",
    r"ignorez?\s+(toutes\s+)?(les\s+)?instructions\s+(précédentes|antérieures)", r"invite\s+système",
    r"exfiltrat", r"send\s+(all|the)\s+(data|documents|files)\s+to\b",
]
INJ_RE = re.compile("|".join(INJECTION_PATTERNS), re.I)


class ExtractionError(Exception):
    pass


@dataclass
class Segment:
    locator: str
    text: str
    injection: bool = False


@dataclass
class Extracted:
    kind: str
    mime: str
    title: str = ""
    author: str = ""
    doc_date: str = ""
    segments: list[Segment] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    ocr_pages: list[int] = field(default_factory=list)  # pages (1-based) with no selectable text

    @property
    def needs_ocr(self) -> bool:
        return bool(self.ocr_pages)


MIME = {"docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "pdf": "application/pdf",
        "txt": "text/plain", "md": "text/markdown", "csv": "text/csv",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "image": "image/*"}
IMAGE_MIME = {b"\x89PNG": "image/png", b"\xff\xd8\xff": "image/jpeg", b"II*\x00": "image/tiff", b"MM\x00*": "image/tiff"}


def sanitize_filename(name: str) -> str:
    name = name.replace("\\", "/").split("/")[-1]
    name = re.sub(r"[^\w.\- ()À-ÿ]", "_", name).strip(" .")
    return (name or "upload")[:150]


def _check_zip(data: bytes) -> zipfile.ZipFile:
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise ExtractionError("Corrupted Office file.") from e
    infos = z.infolist()
    if len(infos) > MAX_ZIP_ENTRIES:
        raise ExtractionError("Archive has too many entries.")
    total = sum(i.file_size for i in infos)
    packed = max(sum(i.compress_size for i in infos), 1)
    if total > MAX_ZIP_UNCOMPRESSED or total / packed > MAX_ZIP_RATIO:
        raise ExtractionError("File expands to an unsafe size (possible archive bomb).")
    return z


def sniff(data: bytes, filename: str) -> str:
    """Identify the real file type from content. Extension alone is never trusted."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if any(data.startswith(m) for m in IMAGE_MIME):
        return "image"
    if data.startswith(b"PK\x03\x04"):
        z = _check_zip(data)
        names = set(z.namelist())
        if "[Content_Types].xml" not in names:
            raise ExtractionError("Unsupported archive.")
        if any(n.startswith("word/") for n in names):
            if any(n.lower().endswith("vbaproject.bin") for n in names):
                raise ExtractionError("Macro-enabled documents are not accepted.")
            return "docx"
        if any(n.startswith("xl/") for n in names):
            if any(n.lower().endswith("vbaproject.bin") for n in names):
                raise ExtractionError("Macro-enabled workbooks are not accepted.")
            return "xlsx"
        raise ExtractionError("Unsupported Office file type.")
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        raise ExtractionError("Legacy binary Office files (.doc/.xls) are not accepted; save as .docx/.xlsx.")
    if b"\x00" in data[:4096]:
        raise ExtractionError("Binary file type not supported.")
    try:
        data[:200000].decode("utf-8")
    except UnicodeDecodeError:
        try:
            data[:200000].decode("cp1252")
        except UnicodeDecodeError as e:
            raise ExtractionError("Text encoding not recognized.") from e
    if ext == "csv":
        return "csv"
    if ext in ("md", "markdown"):
        return "md"
    return "txt"


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _chunk(locator_base: str, paras: list[str]) -> list[Segment]:
    segs, buf, start = [], [], 1
    size = 0
    for i, p in enumerate(paras, 1):
        if size + len(p) > SEG_CHARS and buf:
            segs.append(Segment(f"{locator_base}, ¶ {start}–{i - 1}", "\n".join(buf)))
            buf, size, start = [], 0, i
        buf.append(p)
        size += len(p)
    if buf:
        segs.append(Segment(f"{locator_base}, ¶ {start}–{len(paras)}" if len(paras) > 1 else locator_base, "\n".join(buf)))
    return segs


def _docx(data: bytes, out: Extracted) -> None:
    import docx
    d = docx.Document(io.BytesIO(data))
    cp = d.core_properties
    out.title, out.author = (cp.title or ""), (cp.author or "")
    out.doc_date = cp.modified.date().isoformat() if cp.modified else ""
    heading, paras = "Start", []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith(("heading", "titre", "title")):
            if paras:
                out.segments.extend(_chunk(f"§ {heading}", paras))
            heading, paras = t[:120], []
            if not out.title and style.startswith("title"):
                out.title = t
            continue
        paras.append(t)
    if paras:
        out.segments.extend(_chunk(f"§ {heading}", paras))
    for ti, table in enumerate(d.tables, 1):
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        out.segments.extend(_chunk(f"Table {ti}", [r for r in rows if r.strip(" |")]))


def _pdf(data: bytes, out: Extracted) -> None:
    import pdfplumber
    empty = 0
    empty_pages: list[int] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        meta = pdf.metadata or {}
        out.title, out.author = str(meta.get("Title", "") or ""), str(meta.get("Author", "") or "")
        if len(pdf.pages) > 500:
            raise ExtractionError("PDF has more than 500 pages.")
        for i, page in enumerate(pdf.pages, 1):
            t = (page.extract_text() or "").strip()
            if not t:
                empty += 1
                empty_pages.append(i)
                continue
            paras = [x.strip() for x in re.split(r"\n\s*\n|\n(?=[A-Z0-9][^\n]{0,80}\n)", t) if x.strip()]
            out.segments.extend(_chunk(f"p. {i}", paras))
        out.ocr_pages = empty_pages
        if empty and not out.segments:
            out.warnings.append("No selectable text: this looks like a scanned PDF. Run text recognition (OCR) to read it.")
        elif empty:
            out.warnings.append(f"{empty} page(s) had no selectable text. Run text recognition (OCR) to read them.")


def _text(data: bytes, out: Extracted, md: bool) -> None:
    text = _decode(data)
    heading, paras, line_no = "Start", [], 0
    for block in re.split(r"\n\s*\n", text):
        b = block.strip()
        if not b:
            continue
        if md and b.startswith("#"):
            first, _, rest = b.partition("\n")
            if paras:
                out.segments.extend(_chunk(f"§ {heading}", paras))
            heading, paras = first.lstrip("# ").strip()[:120], []
            if rest.strip():
                paras.append(rest.strip())
            continue
        paras.append(b)
        line_no += 1
    if paras:
        out.segments.extend(_chunk(f"§ {heading}" if md else "Text", paras))


def _csv(data: bytes, out: Extracted) -> None:
    rows = list(csv.reader(io.StringIO(_decode(data))))[:5000]
    lines = [" | ".join(r) for r in rows if any(c.strip() for c in r)]
    for i in range(0, len(lines), 40):
        out.segments.append(Segment(f"rows {i + 1}–{min(i + 40, len(lines))}", "\n".join(lines[i:i + 40])))


def _xlsx(data: bytes, out: Extracted) -> None:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for ws in wb.worksheets[:20]:
        lines = []
        for r_i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if r_i > 5000:
                out.warnings.append(f"Sheet {ws.title}: only the first 5,000 rows were read.")
                break
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append((r_i, " | ".join(cells)))
        for i in range(0, len(lines), 40):
            part = lines[i:i + 40]
            out.segments.append(Segment(f"{ws.title}!rows {part[0][0]}–{part[-1][0]}", "\n".join(x[1] for x in part)))
    wb.close()


def _image(data: bytes, out: Extracted) -> None:
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = 60_000_000  # refuse decompression bombs
    try:
        img = Image.open(io.BytesIO(data))
        img.verify()
        img = Image.open(io.BytesIO(data))
        frames = getattr(img, "n_frames", 1)
    except Image.DecompressionBombError as e:
        raise ExtractionError("Image is too large.") from e
    out.mime = Image.MIME.get(img.format, "image/png")
    out.ocr_pages = list(range(1, min(frames, 100) + 1))
    out.warnings.append("Image file: run text recognition (OCR) to read it.")


def extract(data: bytes, filename: str) -> Extracted:
    kind = sniff(data, filename)
    out = Extracted(kind=kind, mime=MIME[kind])
    try:
        {"docx": _docx, "pdf": _pdf, "csv": _csv, "xlsx": _xlsx, "image": _image}.get(kind, lambda d, o: _text(d, o, kind == "md"))(data, out)
    except ExtractionError:
        raise
    except Exception as e:  # parser failure on hostile or corrupted input
        raise ExtractionError(f"Could not read the file ({type(e).__name__}).") from e
    flagged = 0
    for s in out.segments:
        if INJ_RE.search(s.text):
            s.injection = True
            flagged += 1
    if flagged:
        out.warnings.append(f"{flagged} passage(s) contain text that reads like instructions to an AI system. "
                            "They are treated as document content only and are never followed.")
    if not out.segments and not out.warnings and not out.ocr_pages:
        out.warnings.append("No text could be extracted.")
    return out
